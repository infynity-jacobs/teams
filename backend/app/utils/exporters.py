import io
from pathlib import Path
from typing import List, Sequence, Any, Optional, Mapping

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, portrait, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, Flowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm


class _SVGLogo(Flowable):
    """Render an SVG logo directly as vector graphics in the PDF.

    Using renderPDF avoids the optional raster backends used by renderPM,
    which are not consistently available on Ubuntu installations.
    """
    def __init__(self, drawing, width: float, height: float):
        super().__init__()
        self.drawing = drawing
        self.width = width
        self.height = height
        dw = float(getattr(drawing, "width", 1) or 1)
        dh = float(getattr(drawing, "height", 1) or 1)
        scale = min(width / dw, height / dh)
        self._scale = scale
        self._draw_width = dw * scale
        self._draw_height = dh * scale

    def wrap(self, availWidth, availHeight):
        return self.width, self.height

    def draw(self):
        from reportlab.graphics import renderPDF
        canvas = self.canv
        canvas.saveState()
        canvas.translate(0, (self.height - self._draw_height) / 2.0)
        canvas.scale(self._scale, self._scale)
        renderPDF.draw(self.drawing, canvas, 0, 0)
        canvas.restoreState()


def build_xlsx(headers: Sequence[str], rows: Sequence[Sequence[Any]], title: str = "Report") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else "Report"

    ws.append(list(headers))
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        ws.append(list(row))

    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _pdf_page_layout(headers: Sequence[str], rows: Sequence[Sequence[Any]]):
    """Choose a readable A4 orientation from report width and size."""
    column_count = len(headers)
    row_count = len(rows)
    orientation = landscape if column_count >= 7 or row_count >= 25 else portrait
    page_size = orientation(A4)
    return page_size, orientation is landscape


def _brand_header(branding: Optional[Mapping[str, Any]], title: str, styles, available_width: float):
    branding = branding or {}
    company = str(branding.get("company_name") or branding.get("site_name") or "").strip()
    logo_path = branding.get("logo_path")
    address = str(branding.get("company_address") or "").strip()
    contact = " • ".join(x for x in [
        str(branding.get("contact_phone") or "").strip(),
        str(branding.get("contact_email") or "").strip(),
    ] if x)

    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=18, leading=21,
        spaceAfter=2, alignment=0,
    )
    company_style = ParagraphStyle(
        "ReportCompany", parent=styles["Heading2"], fontSize=12, leading=14,
        textColor=colors.HexColor("#1F4E78"), spaceAfter=2,
    )
    meta_style = ParagraphStyle(
        "ReportMeta", parent=styles["Normal"], fontSize=8, leading=10,
        textColor=colors.HexColor("#666666"),
    )

    text_parts = []
    if company:
        text_parts.append(Paragraph(company, company_style))
    text_parts.append(Paragraph(title, title_style))
    if address or contact:
        details = "<br/>".join(x for x in [address, contact] if x)
        text_parts.append(Paragraph(details, meta_style))

    # Keep the logo local: report generation runs on the server and should not
    # depend on public DNS/HTTPS or the reverse proxy being reachable.
    logo = None
    if logo_path:
        try:
            p = Path(str(logo_path))
            if p.exists() and p.is_file() and p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".svg"}:
                logo_source = str(p)
                if p.suffix.lower() == ".webp":
                    # ReportLab does not reliably decode WebP on all deployments.
                    # Convert it in memory to PNG when Pillow is available.
                    from PIL import Image as PILImage
                    import io as _io
                    with PILImage.open(p) as im:
                        if im.mode not in ("RGB", "RGBA"):
                            im = im.convert("RGBA")
                        converted = _io.BytesIO()
                        im.save(converted, format="PNG")
                        converted.seek(0)
                        logo_source = converted
                elif p.suffix.lower() == ".svg":
                    # Render SVG directly as vector graphics. This avoids relying
                    # on renderPM's optional raster backends (for example
                    # rlPyCairo) which may be absent on a production server.
                    from svglib.svglib import svg2rlg
                    drawing = svg2rlg(str(p))
                    if drawing is None:
                        raise ValueError("Unable to parse SVG logo")
                    logo = _SVGLogo(drawing, width=3.8 * cm, height=1.25 * cm)
                if logo is None:
                    logo = Image(logo_source, width=3.8 * cm, height=1.25 * cm, kind="proportional")
                logo.hAlign = "LEFT"
        except Exception:
            logo = None

    if logo:
        left = [logo]
        right = text_parts or [Paragraph(title, title_style)]
        header = Table([[left, right]], colWidths=[4.4 * cm, max(available_width - 4.4 * cm, 1 * cm)])
        header.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        return header
    return text_parts


def _pdf_footer(canvas, doc, branding: Optional[Mapping[str, Any]]):
    branding = branding or {}
    footer = str(branding.get("report_email_footer") or "").strip()
    company = str(branding.get("company_name") or branding.get("site_name") or "").strip()
    if not footer:
        footer = company
    if not footer:
        return
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#777777"))
    canvas.drawCentredString(doc.pagesize[0] / 2, 0.55 * cm, footer)
    canvas.drawRightString(doc.pagesize[0] - 1.2 * cm, 0.55 * cm, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf(headers: Sequence[str], rows: Sequence[Sequence[Any]], title: str = "Report",
              subtitle: str = "", branding: Optional[Mapping[str, Any]] = None) -> bytes:
    buf = io.BytesIO()
    pagesize, is_landscape = _pdf_page_layout(headers, rows)
    margin = 1.2 * cm
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=1.25 * cm,
    )
    styles = getSampleStyleSheet()
    available_width = pagesize[0] - (2 * margin)
    elements = _brand_header(branding, title, styles, available_width)
    if not isinstance(elements, list):
        elements = [elements]
    if subtitle:
        elements.append(Paragraph(subtitle, styles["Normal"]))
    elements.append(Spacer(1, 0.5 * cm))

    table_data = [list(headers)] + [[str(c) if c is not None else "" for c in row] for row in rows]

    if headers:
        preferred = []
        for idx, header in enumerate(headers):
            values = [str(row[idx]) if idx < len(row) and row[idx] is not None else "" for row in rows[:100]]
            sample_len = max([len(str(header))] + [len(v) for v in values] or [1])
            if header.lower() in {"id", "qty", "leads", "converted", "new", "pending", "lost", "conversion %"}:
                weight = min(max(sample_len, 5), 12)
            elif header.lower() in {"email", "phone", "sku", "created", "converted"}:
                weight = min(max(sample_len, 8), 20)
            else:
                weight = min(max(sample_len, 10), 28)
            preferred.append(weight)
        total = sum(preferred) or 1
        col_widths = [available_width * w / total for w in preferred]
    else:
        col_widths = None

    table = Table(table_data, repeatRows=1, colWidths=col_widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)
    doc.build(elements, onFirstPage=lambda c, d: _pdf_footer(c, d, branding),
              onLaterPages=lambda c, d: _pdf_footer(c, d, branding))
    buf.seek(0)
    return buf.read()
